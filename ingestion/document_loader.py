
from io import BytesIO
from pathlib import Path

import fitz  # PyMuPDF
import pandas as pd
from docx import Document


def extract_text(filename: str, file_bytes: bytes) -> str:
    """
    Extract text from PDF, DOCX, CSV, or TXT files.

    Args:
        filename: Name of the uploaded file.
        file_bytes: File contents as bytes.

    Returns:
        Extracted text as a string.
    """
    extension = Path(filename).suffix.lower()

    if not file_bytes:
        raise ValueError(f"{filename} is empty.")

    try:
        # PDF extraction
        if extension == ".pdf":
            pages = []

            with fitz.open(stream=file_bytes, filetype="pdf") as pdf:
                for page_number, page in enumerate(pdf, start=1):
                    page_text = page.get_text().strip()

                    if page_text:
                        pages.append(
                            f"[Page {page_number}]\n{page_text}"
                        )

            text = "\n\n".join(pages)

        # Word document extraction
        elif extension == ".docx":
            doc = Document(BytesIO(file_bytes))
            parts = []

            for paragraph in doc.paragraphs:
                content = paragraph.text.strip()
                if content:
                    parts.append(content)

            # Extract text from tables too
            for table in doc.tables:
                for row in table.rows:
                    cells = [
                        cell.text.strip()
                        for cell in row.cells
                    ]
                    parts.append(" | ".join(cells))

            text = "\n".join(parts)

        # CSV extraction
        elif extension == ".csv":
            dataframe = pd.read_csv(BytesIO(file_bytes))

            # Convert column names and rows into readable text
            text = dataframe.to_string(index=False)

        # Plain text extraction
        elif extension == ".txt":
            try:
                text = file_bytes.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = file_bytes.decode("latin-1")

        else:
            raise ValueError(
                f"Unsupported file type: {extension}"
            )

    except ValueError:
        raise
    except Exception as error:
        raise ValueError(
            f"Could not read {filename}: {error}"
        ) from error

    text = text.strip()

    if not text:
        raise ValueError(
            f"No readable text was found in {filename}. "
            "If this is a scanned PDF, OCR may be required."
        )

    return text
